import re

from docx import Document

from conference_parser import (
    CODE_PATTERN,
    clean_text,
    make_session,
)


# ============================================================
# DOCX 段落與表格擷取
# ============================================================

def docx_to_tables_and_text(file_path):
    document = Document(str(file_path))

    paragraphs = []

    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)
        if text:
            paragraphs.append(text)

    tables = []

    for table_index, table in enumerate(document.tables, start=1):
        table_rows = []

        for row in table.rows:
            cells = [
                clean_text(cell.text)
                for cell in row.cells
            ]

            if any(cells):
                table_rows.append(cells)

        if table_rows:
            tables.append({
                "table_index": table_index,
                "rows": table_rows,
            })

    return paragraphs, tables


# ============================================================
# DOCX 表格輸出
# ============================================================

def print_docx_tables(tables):
    from conference_parser import log

    for table in tables:
        log("")
        log("=" * 72)
        log(
            f"===== DOCX TABLE "
            f"{table['table_index']} ====="
        )
        log("=" * 72)

        for row_index, row in enumerate(table["rows"], start=1):
            log(
                f"[R{row_index}] "
                + " | ".join(row)
            )


# ============================================================
# DOCX 文字備援解析
# 保留 v11.2 原有邏輯
# ============================================================

def parse_docx_text_sessions(
    text_blocks,
    event_title,
    event_id,
    file_format="docx",
):
    sessions = []

    for block in text_blocks:
        for raw_line in block.splitlines():
            line = clean_text(raw_line)

            if not line:
                continue

            match = CODE_PATTERN.search(line)

            if not match:
                continue

            code = match.group(1)

            rest = clean_text(
                line[:match.start()]
                + " "
                + line[match.end():]
            ).strip(" |｜\t,，:：-")

            if not rest:
                continue

            if any(
                keyword in rest
                for keyword in (
                    "公司名稱",
                    "公司代碼",
                    "股票代碼",
                    "證券代號",
                    "報名",
                    "場地",
                    "時間",
                    "日期",
                )
            ):
                continue

            company_name = re.split(
                r"\s{2,}|[|｜\t]",
                rest,
                maxsplit=1,
            )[0].strip()

            if re.fullmatch(r"\d{1,3}", company_name):
                company_name = ""

            sessions.append(
                make_session(
                    stock_code=code,
                    company_name=company_name,
                    event_title=event_title,
                    event_id=event_id,
                    file_format=file_format,
                    parse_method="text_fallback",
                    raw_text=line,
                )
            )

    return sessions
